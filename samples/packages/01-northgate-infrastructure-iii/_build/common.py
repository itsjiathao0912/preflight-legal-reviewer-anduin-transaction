"""Shared python-docx helpers for the Northgate sample package (all text is fictional)."""
import re
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

BANNER = "FICTIONAL SAMPLE — test data for Sub-doc QA — not a real offering"


def _field(run, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def new_doc(title):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        h = d.styles[name]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = None
    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    r = hp.add_run(BANNER + "\n" + title)
    r.font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(title + " — page ")
    r.font.size = Pt(8)
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    _field(r2, "PAGE")
    d.core_properties.title = title
    d.core_properties.author = "Fictional sample (Sub-doc QA test data)"
    d.core_properties.comments = BANNER
    return d


_SUP = re.compile(r"\^(\d+)\^")


def add_text(p, text):
    """Add text to paragraph; ^1^ becomes a superscript footnote marker; **x** bold."""
    parts = re.split(r"(\^\d+\^|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        m = _SUP.fullmatch(part)
        if m:
            run = p.add_run(m.group(1))
            run.font.superscript = True
        elif part.startswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)
    return p


def para(d, text, bold=False, indent=0.0, italic=False, size=None):
    p = d.add_paragraph()
    add_text(p, text)
    if bold or italic or size:
        for r in p.runs:
            r.bold = r.bold or bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
    if indent:
        p.paragraph_format.left_indent = Pt(indent * 18)
    return p


def _hd(d, t, level):
    p = d.add_heading(t, level=level)
    for r in p.runs:
        r.bold = True  # explicit run bold so HTML-converting readers see a bold cue
    return p


def h1(d, t):
    return _hd(d, t, 1)


def h2(d, t):
    return _hd(d, t, 2)


def h3(d, t):
    return _hd(d, t, 3)


def table(d, rows, header=True, widths=None):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            add_text(c.paragraphs[0], str(cell))
            if header and i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True
    d.add_paragraph()
    return t


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


class Footnotes:
    """Collect footnotes for a section; flush prints a footnote block."""

    def __init__(self, d):
        self.d = d
        self.items = []

    def mark(self, text):
        self.items.append(text)
        return f"^{len(self.items)}^"

    def flush(self):
        if not self.items:
            return
        para(self.d, "Footnotes", bold=True, size=9)
        for i, t in enumerate(self.items, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"{i}. {t}")
            r.font.size = Pt(9)
        self.items = []


def words_of_docx(path):
    d = Document(path)
    out = []
    for p in d.paragraphs:
        out.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    return "\n".join(out)
