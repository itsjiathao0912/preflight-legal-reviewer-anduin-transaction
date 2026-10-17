#!/usr/bin/env python3
"""Check the generated package: word counts, pages, headings, seeded quotes, internal consistency."""
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

QUOTES = {
    "F1": "If you answered Yes to Question 6 (Benefit Plan Investor), skip to Question 15.",
    "F2a": "Annex B (ERISA Representations) must be completed by each Investor that is an ERISA plan, but is not required for plans subject to Section 4975 of the Code (e.g., IRAs)",
    "F3": "the 25% limitation described in Section 8.4",
    "F4a": "Plan Assets",
    "F4b": "Significant Participation",
    "F4c": "Hurdle Amount",
    "F4d": "Subscription Facility",
    "F5": "Qualified Credit Investor",
    "F6a": "Taxable U.S. person",
    "F6b": "U.S. person subject to tax",
    "F7": "[NTD: SL&A to confirm whether the foundation needs a separate UBTI rep.]",
    "F8": "[insert Subscription Facility lender]",
    "F9": "Commitments may be reduced from time to time at the discretion of the General Partner or the Investor",
    "F10a": "ABA routing number: 999000011",
    "F10b": "Account number: 9900123456",
}


def full_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    ok = True
    print(f"{'file':44} {'words':>6} {'pages~':>7} {'headings':>8} {'tables':>6} {'images':>6}")
    texts = {}
    for f in sorted(DOCS.glob("*.docx")):
        d = Document(f)
        txt = full_text(d)
        texts[f.name] = (d, txt)
        words = len(txt.split())
        heads = sum(1 for p in d.paragraphs if p.style.name.startswith("Heading"))
        imgs = len(d.inline_shapes)
        print(f"{f.name:44} {words:6d} {words / 475:7.1f} {heads:8d} {len(d.tables):6d} {imgs:6d}")
    sub_name = "01-RR1-Subscription-Agreement.docx"
    sub_doc, sub = texts[sub_name]
    wc = len(sub.split())
    print(f"\nsub agreement >= 7,500 words: {wc >= 7500} ({wc}); in 8,000-11,000: {8000 <= wc <= 11000}")
    ok &= wc >= 7500

    bounds = {"02-RR1-LPA-Definitions.docx": (3000, 4500), "03-RR1-PPM-Summary.docx": (1800, 2500),
              "04-Pinecroft-AML-KYC-Checklist.docx": (1000, 1500), "05-GP-Requirements-Email.docx": (400, 700)}
    for k, (lo, hi) in bounds.items():
        w = len(texts[k][1].split())
        flag = lo <= w <= hi
        ok &= flag
        print(f"  {k}: {w} in [{lo},{hi}] -> {flag}")
    scan_doc, scan_txt = texts["06-RR1-Tax-Forms-Guide-SCANNED.docx"]
    print(f"  scan: body text chars={len(scan_txt.strip())} (must be 0), images={len(scan_doc.inline_shapes)} (must be 1)")
    ok &= len(scan_txt.strip()) == 0 and len(scan_doc.inline_shapes) == 1

    print("\nSeeded quotes in 01 (verbatim):")
    for k, q in QUOTES.items():
        found = q in sub
        count = sub.count(q)
        print(f"  {k:5} found={found} x{count}  {q[:70]}")
        ok &= found

    # F2 footnote markers: superscript '3' on 401(k), DB, IRA rows; none on "Other benefit plan"
    q3 = [t for t in sub_doc.tables if any("Other benefit plan" in c.text for r in t.rows for c in r.cells)][0]
    print("\nF2 footnote markers (Question 3 table):")
    for r in q3.rows[1:]:
        cell = r.cells[1]
        sup = "".join(run.text for p in cell.paragraphs for run in p.runs if run.font.superscript)
        print(f"  {cell.text[:62]:62} marker={sup or '-'}")

    # F4: LPA defines the four terms; sub agreement does not define them
    lpa = texts["02-RR1-LPA-Definitions.docx"][1]
    print("\nF4/F5 definitions:")
    for term in ("Plan Assets", "Significant Participation", "Hurdle Amount", "Subscription Facility"):
        in_lpa = f"“{term}” means" in lpa
        in_sub = f"“{term}” means" in sub
        print(f"  {term:28} defined in LPA={in_lpa}  defined in sub={in_sub}")
        ok &= in_lpa and not in_sub
    qci = any("Qualified Credit Investor" in t for _, t in texts.values() if t is not texts[sub_name][1]
              ) or ("“Qualified Credit Investor”" in sub)
    print(f"  Qualified Credit Investor defined anywhere: {qci} (must be False)")
    ok &= not qci

    # Cross-reference integrity of the sub agreement (only F3 may be wrong)
    secs = set(re.findall(r"^\*{0,2}(\d+\.\d+) ", "\n".join(p.text for p in sub_doc.paragraphs), re.M))
    heads = {m for p in sub_doc.paragraphs if p.style.name == "Heading 1"
             for m in re.findall(r"^Section (\d+)\.", p.text)}
    refs = re.findall(r"Section (\d+(?:\.\d+)?)(?:\(\w+\))?( of (?:the )?(?:Code|ERISA|Securities Act|Investment Company Act))?", sub)
    bad = sorted({r for r, ext in refs if not ext and (r not in secs and r not in heads)})
    print(f"\nInternal Section refs unresolved: {bad} (must be [])")
    ok &= not bad
    qs = {int(x) for x in re.findall(r"\*?\*?Question (\d+)\.\*\*", "\n".join(p.text for p in sub_doc.paragraphs))}
    qs |= {int(x) for x in re.findall(r"^Question (\d+)\.", "\n".join(p.text for p in sub_doc.paragraphs), re.M)}
    qrefs = sorted({int(x) for x in re.findall(r"Question (\d+)", sub)})
    print(f"Questions present: {min(qs)}-{max(qs)} ({len(qs)}); referenced: {qrefs}; unresolved: {[q for q in qrefs if q not in qs]}"
          " (only 15 expected)")
    print("Brackets in 01:", re.findall(r"\[[^\]]{1,80}\]", sub))
    print("Digits cued by routing/account in 01:", re.findall(r"(?i)(?:routing|account|ABA)[^\n]{0,30}?\d{6,}", sub))
    for a in ("Annex A", "Annex B", "Annex C", "Annex D"):
        print(f"  {a} heading present:", any(p.style.name == 'Heading 1' and p.text.startswith(a) for p in sub_doc.paragraphs))
    ok_msg = "ALL CHECKS PASSED" if ok else "SOME CHECKS FAILED"
    print("\n" + ok_msg)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
