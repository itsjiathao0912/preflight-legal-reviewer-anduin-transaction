#!/usr/bin/env python3
"""Check script: word counts, estimated pages, headings, seeded quotes, internal reference sanity."""
import glob, os, re, sys, collections
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(HERE, "..", "docs")


def texts(path):
    d = Document(path)
    out = []
    for p in d.paragraphs:
        out.append(p.text)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                out.append(c.text)
    return d, out


QUOTES = {
    "E1a Capital Commitment": "Capital Commitment",
    "E1b Drawdown Notice": "Drawdown Notice",
    "E1c Defaulting Limited Partner": "Defaulting Limited Partner",
    "E1d Excused Investor": "Excused Investor",
    "E1e Key Person Event": "Key Person Event",
    "E1f Carried Interest Escrow": "Carried Interest Escrow",
    "E1g LPA fallback sentence": "Capitalized terms used but not defined herein have the meanings given to them in the Partnership Agreement.",
    "E2 Annex C ref": "complete the Cayman AML Schedule set out in Annex C",
    "E3a Sec 3.2": "All Investors must deliver the documents listed in Annex C.",
    "E3b Sec 9.5": "Investors that are regulated financial institutions in a FATF member jurisdiction are not required to deliver any documents under Annex C.",
    "E4 Sec 5.1": "accredited investor or a professional investor, and the General Partner and the Fund may treat those two expressions as interchangeable",
    "E5a correct transfer clause": "the General Partner shall not unreasonably withhold consent",
    "E5b typo transfer clause": "the General Partner shall unreasonably withhold consent",
    "E7 S$": "For Investors resident in Singapore, the minimum Capital Commitment is S$3,000,000.",
    "E8 placeholder": "[●]% of the carried interest",
    "E9a account": "Account Number: 9990001234567",
    "E9b routing": "ABA Routing No.: 000000018",
}

print("== Word counts ==")
files = sorted(glob.glob(os.path.join(DOCS, "*.docx")))
info = {}
for f in files:
    d, tx = texts(f)
    words = sum(len(t.split()) for t in tx)
    heads = sum(1 for p in d.paragraphs if p.style.name.startswith("Heading"))
    h1 = sum(1 for p in d.paragraphs if p.style.name == "Heading 1")
    hdr = d.sections[0].header.paragraphs[0].text
    info[os.path.basename(f)] = (words, tx)
    print(f"{os.path.basename(f):58s} words={words:6d} est_pages={words/475:5.1f} headings={heads:3d} (H1={h1:2d}) header_ok={'FICTIONAL SAMPLE' in hdr}")

v2 = [k for k in info if k.startswith("01-")][0]
v1 = [k for k in info if k.startswith("04-")][0]
ok = info[v2][0] >= 7500
print(f"\nMain sub agreement >= 7,500 words: {ok} ({info[v2][0]})")
targets = {"01-": (8000, 12000), "02-": (2500, 3500), "03-": (2000, 3000), "04-": (7500, 10000), "05-": (800, 1200), "06-": (900, 1300)}
for k, (lo, hi) in targets.items():
    n = [x for x in info if x.startswith(k)][0]
    w = info[n][0]
    print(f"  {n[:2]} range {lo}-{hi}: {'OK' if lo <= w <= hi else 'OUT OF RANGE'} ({w})")

print("\n== Seeded quotes in 01 (v2) ==")
blob2 = "\n".join(info[v2][1])
for k, q in QUOTES.items():
    print(f"  {'PRESENT' if q in blob2 else 'MISSING':8s} x{blob2.count(q):2d}  {k}")
print("\n== Same quotes in 04 (v1 baseline) ==")
blob1 = "\n".join(info[v1][1])
for k in ("E2 Annex C ref", "E3a Sec 3.2", "E3b Sec 9.5", "E5a correct transfer clause", "E5b typo transfer clause", "E7 S$", "E8 placeholder", "E9a account"):
    print(f"  x{blob1.count(QUOTES[k]):2d}  {k}")
print("  v1 minimum / escrow:", "US$5,000,000" in blob1, "25% of the carried interest" in blob1, "Vietnam" in blob1, "Politically Exposed Persons Questionnaire" in blob1)
print("  v2 minimum / escrow:", "US$3,000,000" in blob2, "30% of the carried interest" in blob2, "Vietnam" in blob2, "Politically Exposed Persons Questionnaire" in blob2)

print("\n== Extra-defect guards in 01 (v2) ==")
print("  '[' placeholders :", re.findall(r"\[[^\]]{0,20}\]", blob2))
print("  S$ occurrences (excluding US$):", len(re.findall(r"(?<!U)S\$", blob2)))
print("  'Annex C' count  :", blob2.count("Annex C"))
print("  Drafting notes   :", re.findall(r"NTD|Drafting note|Note to draft", blob2))
print("  Real-world names :", [n for n in ("Linklaters", "Clifford", "Citco", "Deloitte", "KPMG", "Ernst", "PwC", "Allen & Overy", "HSBC", "Goldman") if n in blob2])

print("\n== Internal 'Section N.M' references resolve (01 v2 and 04 v1) ==")
for name in (v2, v1):
    d, tx = texts(os.path.join(DOCS, name))
    heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    secs = {int(m.group(1)) for h in heads if (m := re.match(r"Section (\d+)\.", h))}
    clauses = set()
    for p in d.paragraphs:
        m = re.match(r"(\d+\.\d+) ", p.text)
        if m:
            clauses.add(m.group(1))
    annex_h = {m.group(1) for h in heads if (m := re.match(r"Annex ([A-Z])\.", h))}
    bad = []
    for t in "\n".join(tx).split("\n"):
        for m in re.finditer(r"Section (\d+(?:\.\d+)?)", t):
            ref = m.group(1)
            if "." in ref:
                if ref not in clauses:
                    bad.append(ref)
            elif int(ref) not in secs:
                bad.append(ref)
    annex_refs = sorted(set(re.findall(r"Annex ([A-Z])\b", "\n".join(tx))))
    print(f"  {name[:2]}: sections={sorted(secs)[0]}-{sorted(secs)[-1]}  broken Section refs={sorted(set(bad))}  annex headings={sorted(annex_h)} annex refs={annex_refs}")
    d_part = {p.text for p in d.paragraphs if p.style.name == "Heading 2" and p.text.startswith("Part")}
    print(f"      Annex A parts: {sorted(d_part)}")
    missing_annex = sorted(set(annex_refs) - annex_h)
    print(f"      Annex letters referenced but not present: {missing_annex}")

print("\n== Capitalised multi-word terms in 01 not defined in 1.1 (review list) ==")
d, tx = texts(os.path.join(DOCS, v2))
body = "\n".join(p.text for p in d.paragraphs)
defined = set(re.findall(r"“([A-Z][A-Za-z ]+)”", body))
cands = collections.Counter(m for m in re.findall(r"(?<![.\n] )\b([A-Z][a-z]+(?: [A-Z][a-z]+)+)\b", body))
for t, n in sorted(cands.items()):
    if t not in defined:
        print(f"  {n:3d}  {t}")
